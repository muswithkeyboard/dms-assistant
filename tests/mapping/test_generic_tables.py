"""Таблицы документов как пункты запасного адаптера и программа «*» (этап w10-fix, дополнение CONTRACTS §11).

Свои маленькие docx (названия и тексты выдуманы): строка таблицы — пункт «<документ> › табл. N, стр. M» с ячейками
через « | »; текстовая рамка (одна ячейка — много абзацев) — абзацы-пункты с разделами по заголовкам; guided читает
строки так же, как строки таблиц адаптеров; строки размеченного перечня, ставшие строками перечня, — не пункты,
остальные — с программой разметки; «*» — все программы договора. На комплекте (kit_or_skip) c06 — метки с выбором
клиники совпадают с исходным договором чаще, чем до этапа (19/36), ложных «можно» / «нельзя» против эталона нет."""
from __future__ import annotations

import shutil
import sys
from pathlib import Path

from dms_core import guided, onboarding
from dms_core import table_mapping as tm
from dms_core.adapters import generic
from dms_core.kit import KitContract, KitFile
from dms_core.schemas import ClinicRow, TableMapping, TableRef
from fixtures import client_kit as ck
from fixtures import mapping_fixture as mf
from helpers import ROOT, kit_or_skip, skip

FRAME = ["Программа «Стандарт»", "Входит в программу:", "1. Приемы врачей-специалистов.",
         "1.1. Консультации аллерголога.", "Исключения из программы:", "1. Не оплачиваются услуги:",
         "1.1. Косметология и пластическая хирургия."]


def _docx(path: Path) -> Path:
    """Договор: пункт, таблица услуг (с пустой строкой и объединённой строкой-разделом) и текстовая рамка."""
    from docx import Document

    d = Document()
    d.add_paragraph("1. Предмет договора")
    t = d.add_table(rows=0, cols=3)
    for item in [["Услуга", "Входит", "Комментарий"], ["Аллерголог", "да", ""], ["", "", ""],
                 "SECTION:Стационар", ["2.1", "Плановая госпитализация", "нет"]]:
        row = t.add_row()
        if isinstance(item, str):
            merged = row.cells[0].merge(row.cells[-1])
            merged.text = item.split(":", 1)[1]
        else:
            for c, v in zip(row.cells, item, strict=True):
                c.text = v
    d.add_paragraph("2. Программа")
    frame = d.add_table(rows=1, cols=1)
    cell = frame.rows[0].cells[0]
    cell.text = FRAME[0]
    for line in FRAME[1:]:
        cell.add_paragraph(line)
    path.parent.mkdir(parents=True, exist_ok=True)
    d.save(path)
    return path


def _kc(p: Path) -> KitContract:
    f = KitFile(path=p.name, title="Договор", fmt="docx", doc_type="contract", scope="contract",
                size=p.stat().st_size, sha256="0")
    return KitContract(id="u01", group_id="ug1", legal_entity="ООО", files=[f])


def _parse():
    tmp = ck.workdir()
    p = _docx(tmp / "Договор.docx")
    return generic.parse(_kc(p), tmp)


def test_generic_table_rows_become_clauses():
    out = _parse()
    by_id = {c.id: c for c in out.clauses}
    doc = out.documents[0].id
    assert by_id[f"{doc}#p0"].number == "1"                          # абзац тела — как раньше
    head, row = by_id[f"{doc}#t0r0"], by_id[f"{doc}#t0r1"]
    assert head.path == "Договор › табл. 1, стр. 1" and head.text == "Услуга | Входит | Комментарий"
    assert row.text == "Аллерголог | да" and row.number is None and row.program is None
    assert f"{doc}#t0r2" not in by_id                                 # пустая строка — не пункт
    assert by_id[f"{doc}#t0r3"].text == "Стационар"                   # объединённая строка — без повторов
    assert by_id[f"{doc}#t0r4"].number == "2.1"                       # первая ячейка — номер пункта
    # порядок тела: абзац «2. Программа» после строк таблицы, рамка — после него
    ids = [c.id for c in out.clauses]
    assert ids.index(f"{doc}#t0r4") < ids.index(f"{doc}#p1") < ids.index(f"{doc}#t1r0p0")


def test_text_frame_paragraphs_get_sections_by_headings():
    out = _parse()
    doc = out.documents[0].id
    lines = [c for c in out.clauses if c.id.startswith(f"{doc}#t1r0p")]
    assert [c.text for c in lines] == [x.split(" ", 1)[1] if x[0].isdigit() else x for x in FRAME]
    inc, exc = lines[3], lines[6]
    assert inc.number == "1.1" and exc.number == "1.1"
    # нумерация началась заново под «Исключения…» — другой раздел пути, родитель «1» — свой
    assert inc.path == "Договор › таблица 2 › Входит в программу › 1.1"
    assert exc.path == "Договор › таблица 2 › Исключения из программы › 1.1"
    assert not guided._is_table_row(inc)


def test_guided_reads_generic_rows_like_adapter_rows():
    out = _parse()
    doc = out.documents[0].id
    secs = guided._Sections(out.clauses)
    by_id = {c.id: c for c in out.clauses}
    assert guided._is_table_row(by_id[f"{doc}#t0r1"])
    assert secs.section(by_id[f"{doc}#t0r1"])[0] == "included"        # ячейка «да»
    assert secs.section(by_id[f"{doc}#t0r4"])[0] == "negative"        # ячейка «нет»
    frame = [c for c in out.clauses if c.id.startswith(f"{doc}#t1r0p")]
    assert secs.section(frame[3])[0] == "included"
    assert secs.section(frame[6])[0] == "exclusion"


def test_xlsx_rows_are_table_clauses_with_row_ref_ids():
    tmp = ck.workdir()
    p = mf.clinics_xlsx(tmp / "Перечень.xlsx")
    f = KitFile(path=p.name, title="Перечень", fmt="xlsx", doc_type="clinics", scope="contract",
                size=p.stat().st_size, sha256="0")
    out = generic.parse(KitContract(id="u01", group_id="ug1", legal_entity="ООО", files=[f]), tmp)
    doc = out.documents[0].id
    cl = {c.id: c for c in out.clauses}
    ref = TableRef(doc_id=doc, kind="xlsx_sheet", sheet="Список МО")
    assert tm.row_ref(ref, 4) in cl and cl[tm.row_ref(ref, 4)].path == "Перечень › табл. 1, стр. 5"
    assert cl[tm.row_ref(ref, 4)].text.startswith("Стандарт | Клиника Ромашка")


def _setup(tmp, src: Path):
    reg = ck.registry(tmp)
    g = reg.create_client_group("Клиент Т")
    cc = reg.add_client_contract(g.id, "ООО Ромашка")
    ck.put_file(reg, ck.clients_root(tmp), g.id, cc.id, src, doc_type="contract")
    return reg, cc


def _tables(reg, st, tmp, cc):
    kc = onboarding.kit_contract(reg.client_group(cc.group_id), cc, reg.client_files(cc.id))
    return tm.list_tables(kc, ck.clients_root(tmp), {d.source_path: d.id for d in st.documents(cc.id)})


def _process(reg, st, tmp, cc):
    return onboarding.process(cc.id, store=st, registry=reg, root=ck.clients_root(tmp), run_selfcheck=False)


def test_mapped_list_rows_leave_clauses_and_others_get_program():
    tmp = ck.workdir()
    reg, cc = _setup(tmp, mf.program_tables_docx(tmp / "src" / "Программа.docx"))
    st = ck.kit_store(tmp)
    _process(reg, st, tmp, cc)
    tables = _tables(reg, st, tmp, cc)
    rows_before = {c.id for c in st.clauses(cc.id) if "#t" in c.id}
    assert rows_before                                               # строки таблиц — пункты
    m = TableMapping(header_row=None, columns={0: "clinic_name", 2: "address", 3: "code"}, program_from_title=True)
    for ref, _t, _r in tables:
        reg.save_table_mapping(cc.id, ref, m)
    rep = _process(reg, st, tmp, cc)
    assert rep.counts["clinics"] == 6
    clauses = {c.id: c for c in st.clauses(cc.id)}
    sources = {r.source_ref for r in st.clinics(cc.id)}
    assert not sources & set(clauses)                                # строка перечня — не пункт
    doc = tables[0][0].doc_id
    assert clauses[f"{doc}#t0r0"].program == "Бизнес" and clauses[f"{doc}#t0r1"].program == "Бизнес"
    assert clauses[f"{doc}#t1r0"].program == "Тверь" and clauses[f"{doc}#t1r1"].text == "Поликлиника"
    assert clauses[f"{doc}#t1r1"].program == "Тверь"


def _programs_with_text_docx(path: Path) -> Path:
    """Перечни по таблице на программу, после каждого — таблица-рамка с текстом программы (как у «Ренессанса»)."""
    from docx import Document

    d = Document()
    d.add_paragraph("Приложение 1. Программа ДМС")
    for title, rows in mf.PROGRAM_TABLES.items():
        t = d.add_table(rows=0, cols=4)
        for item in [f"SECTION:{title}", *rows]:
            row = t.add_row()
            if isinstance(item, str):
                merged = row.cells[0].merge(row.cells[-1])
                merged.text = item.split(":", 1)[1]
            else:
                for c, v in zip(row.cells, item, strict=True):
                    c.text = v
        frame = d.add_table(rows=1, cols=1).rows[0].cells[0]
        frame.text = f"Программа «{title}»"
        for line in ["Исключения из программы:", f"1. Не оплачивается услуга {title}."]:
            frame.add_paragraph(line)
    path.parent.mkdir(parents=True, exist_ok=True)
    d.save(path)
    return path


def test_text_after_titled_list_gets_its_program():
    """Таблица-рамка с текстом программы сразу после перечня «программа — из названия таблицы» — пункты этой
    программы (а не общие: иначе вопрос об одной программе читал бы исключения другой)."""
    tmp = ck.workdir()
    reg, cc = _setup(tmp, _programs_with_text_docx(tmp / "src" / "Программа.docx"))
    st = ck.kit_store(tmp)
    _process(reg, st, tmp, cc)
    m = TableMapping(header_row=None, columns={0: "clinic_name", 2: "address", 3: "code"}, program_from_title=True)
    lists = [t for t in _tables(reg, st, tmp, cc) if tm.n_columns(t[2]) == 4]
    assert len(lists) == 2
    for ref, _t, _r in lists:
        reg.save_table_mapping(cc.id, ref, m)
    assert _process(reg, st, tmp, cc).counts["clinics"] == 6
    doc = lists[0][0].doc_id
    frames = {c.id: c.program for c in st.clauses(cc.id) if "r0p" in c.id}
    assert frames[f"{doc}#t1r0p2"] == "Бизнес" and frames[f"{doc}#t3r0p2"] == "Тверь"


def test_all_programs_long_matrix_and_list():
    ref = TableRef(doc_id="u01.contract", kind="xlsx_sheet", sheet="Услуги")
    rows = [["Вид помощи", "Входит"], ["Стоматологическая помощь", "да"], ["Помощь на дому", "нет"]]
    m = TableMapping(role="care_matrix", matrix_layout="long", header_row=0, columns={0: "care_type", 1: "present"},
                     program=tm.ALL_PROGRAMS)
    res = tm.apply_matrix_rows("u01", ref, rows, m, ["Стандарт", "Премиум"])
    assert sorted((x.program, x.care_type, x.present) for x in res.rows) == [
        ("Премиум", "Помощь на дому", False), ("Премиум", "Стоматологическая помощь", True),
        ("Стандарт", "Помощь на дому", False), ("Стандарт", "Стоматологическая помощь", True)]
    assert not res.unknown
    none = tm.apply_matrix_rows("u01", ref, rows, m, [])
    assert not none.rows and none.no_programs
    assert any("все программы договора" in w for w in tm.matrix_warnings("Лист", none))
    r = ClinicRow(id="u01.contract#t1r2", contract_id="u01", doc_id="u01.contract", program=tm.ALL_PROGRAMS,
                  care_type="Поликлиника", clinic_name="Клиника", source_ref="u01.contract#t1r2")
    got = tm.expand_all_programs([r], ["Стандарт", "Премиум"])
    assert [(x.id, x.program, x.source_ref) for x in got] == [
        ("u01.contract#t1r2#v1", "Стандарт", "u01.contract#t1r2"), ("u01.contract#t1r2#v2", "Премиум",
                                                                    "u01.contract#t1r2")]
    assert tm.expand_all_programs([r], []) == []


def test_all_programs_list_in_process():
    """Перечень с колонкой программы даёт программы; второй перечень «для всех программ» — строки каждой."""
    tmp = ck.workdir()
    reg, cc = _setup(tmp, mf.matrix_docx(tmp / "src" / "Договор.docx"))
    st = ck.kit_store(tmp)
    _process(reg, st, tmp, cc)
    (lst, _t1, _r1), (wide, _t2, _r2) = _tables(reg, st, tmp, cc)
    reg.save_table_mapping(cc.id, lst, TableMapping(header_row=0, columns={0: "program", 1: "clinic_name",
                                                                           2: "address", 3: "care_type"}))
    # таблица видов помощи — как «перечень для всех программ» (название — колонка вида помощи; проверка механизма)
    reg.save_table_mapping(cc.id, wide, TableMapping(header_row=0, columns={0: "clinic_name"},
                                                     program=tm.ALL_PROGRAMS))
    rep = _process(reg, st, tmp, cc)
    assert rep.counts["programs"] == 2
    rows = [r for r in st.clinics(cc.id) if r.source_ref.startswith(wide.doc_id + f"#t{wide.index}")]
    assert rows and {r.program for r in rows} == {"Стандарт", "Премиум"}
    assert all(r.program != tm.ALL_PROGRAMS for r in st.clinics(cc.id))
    assert "*" not in [p.name for p in st.contract(cc.id).programs]


def test_c06_labels_with_pick_match_source_more_often():
    kit_or_skip()
    from dms_core.config import derived_dir

    src_db = derived_dir() / "kit.sqlite"
    if not src_db.is_file():
        skip("нет data/derived/kit.sqlite")
    sys.path.insert(0, str(ROOT / "scripts" / "bench"))
    import mapping_compare as mc

    if not mc._questions("c06"):
        skip("нет data/derived/review/review6.jsonl, review7.jsonl")
    from helpers import env

    tmp = ck.workdir()
    (tmp / "derived").mkdir(parents=True)
    shutil.copyfile(src_db, tmp / "derived" / "kit.sqlite")          # копия: общая база не трогается
    from dms_core.store import KitStore

    st = KitStore(tmp / "derived" / "kit.sqlite")
    reg = ck.registry(tmp)
    croot = ck.clients_root(tmp)
    cc = mc.load_client(reg, st, croot, "c06")
    onboarding.process(cc.id, store=st, registry=reg, root=croot, run_selfcheck=False)
    programs = [p.name for p in st.contract(cc.id).programs]
    for ref, m in mc.operator_mappings(mc.client_tables(reg, st, croot, cc), programs, derive=True):
        reg.save_table_mapping(cc.id, ref, m)
    rep = onboarding.process(cc.id, store=st, registry=reg, root=croot, run_selfcheck=False)
    assert rep.readiness == "ready"
    with env(DMS_ENGINE_MODE="guided"):
        t = mc.compare_answers(st, "c06", cc.id, None)
    assert t["label_equal_picked"] > 19, t
    for k in ("false_allowed_gold", "false_denied_gold", "false_allowed_gold_nopick", "false_denied_gold_nopick"):
        assert t[k] == 0, (k, t)
    shutil.rmtree(tmp, ignore_errors=True)
