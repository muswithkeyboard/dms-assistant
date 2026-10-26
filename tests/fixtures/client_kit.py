"""Подключение клиента в тестах (ADR 0011): своя папка данных, реестр (AppStore), файлы клиента.

Файлы клиента — свои маленькие docx (ren_fixture, generic_docx) или копии договора комплекта во временной папке;
ничего не пишется в общую data/ (правило данных этапа w8)."""
from __future__ import annotations

import datetime as dt
import hashlib
import shutil
import tempfile
from pathlib import Path

from dms_api.store import AppStore
from dms_core.kit import classify
from dms_core.schemas import ClientFile
from dms_core.store import KitStore
from fixtures import ren_fixture


def workdir() -> Path:
    return Path(tempfile.mkdtemp(prefix="dms-onb-"))


def registry(tmp: Path) -> AppStore:
    return AppStore(tmp / "app" / "app.sqlite")


def kit_store(tmp: Path) -> KitStore:
    return KitStore(tmp / "derived" / "kit.sqlite")


def clients_root(tmp: Path) -> Path:
    return tmp / "derived" / "clients"


def put_file(reg: AppStore, root: Path, group_id: str, contract_id: str, src: Path, *, name: str | None = None,
             doc_type: str | None = None, scope: str = "contract") -> ClientFile:
    """Копия src в папку клиента и запись в реестр (как после загрузки через API)."""
    name = name or src.name
    owner, folder = (group_id, "_group") if scope == "group" else (contract_id, contract_id)
    dst = root / group_id / folder / name
    dst.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(src, dst)
    data = dst.read_bytes()
    f = ClientFile(name=name, doc_type=doc_type or classify(name)[0], scope=scope, fmt=src.suffix.lstrip(".").lower(),
                   size=len(data), sha256=hashlib.sha256(data).hexdigest(), uploaded_at=dt.datetime.now(dt.UTC))
    reg.save_client_file(owner, f)
    return f


def ren_files(tmp: Path) -> list[Path]:
    """Свой маленький договор «Ренессанса» (tests/fixtures/ren_fixture.py): договор + программа с перечнем."""
    kc, root = ren_fixture.build(tmp / "src")
    return [root / f.path for f in kc.files]


def generic_docx(path: Path, paragraphs: list[str]) -> Path:
    """Договор неизвестной СК: только абзацы текста (адаптер generic)."""
    from docx import Document

    d = Document()
    for t in paragraphs:
        d.add_paragraph(t)
    path.parent.mkdir(parents=True, exist_ok=True)
    d.save(path)
    return path


def rename_clinic(program: Path, old: str, new: str) -> None:
    """Перечень клиник программы с другой клиникой: название в ячейке таблицы заменяется (повторная обработка)."""
    from docx import Document

    d = Document(program)
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                if cell.text.startswith(old):
                    cell.text = cell.text.replace(old, new)
    d.save(program)


def simple_pdf(path: Path, text: str | None) -> Path:
    """Минимальный PDF с одной страницей (текст Helvetica — латиница; None — страница без текстового слоя)."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode() if text else b""
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            (b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
             b"/Resources << /Font << /F1 5 0 R >> >> >>"),
            b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, o in enumerate(objs, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + o + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    out += b"".join(b"%010d 00000 n \n" % off for off in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(bytes(out))
    return path


def simple_xlsx(path: Path, rows: list[list[str | None]]) -> Path:
    """Книга с одним листом (своя маленькая таблица)."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def put_bytes(reg: AppStore, root: Path, group_id: str, contract_id: str, name: str, data: bytes,
              *, doc_type: str = "other", scope: str = "contract") -> ClientFile:
    """Файл с заданным содержимым (битый файл) — в папку клиента и в реестр."""
    src = Path(tempfile.mkdtemp(prefix="dms-onb-src-")) / name
    src.write_bytes(data)
    return put_file(reg, root, group_id, contract_id, src, doc_type=doc_type, scope=scope)
