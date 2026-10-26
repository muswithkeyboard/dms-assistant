#!/usr/bin/env python3
"""Импорт эталонов организаторов К10–К19 → data/derived/gold/k10_k19.jsonl (+ acceptance_k10_k19.json).

ADR 0008: эталоны — производное от документов организаторов, поэтому лежат вне git (data/derived/gold/).
Схема записи — eval/gold/README.md.

Источники:
  * xlsx организаторов (config.gold_xlsx(), лист «Кейсы») — обращение, вопрос, эталонный ответ, вывод SC 01,
    «Документы и пункты», «Что проверяет кейс»;
  * ручная разметка data/derived/gold/k10_k19_manual.yaml (тоже вне git) — сценарий, контекст, ключевые факты,
    обязательные ссылки с ТОЧНЫМ номером пункта по документу и номером из эталона (clause_gold), ловушки.

Каждая ссылка разметки проверяется по самим документам комплекта: блок (абзац pN / строка таблицы tNrM / строка xlsx)
должен содержать quote_hint и всё из also, а номер автонумерации абзаца — совпадать с clause. Хоть одна
непроверенная ссылка → код возврата 1 и файлы не пишутся (эталон с битой ссылкой хуже, чем никакого).
В журнал не печатаются тексты документов и цитаты — только ID кейсов и адреса блоков (`docs/rules.md` п. 4).

Запуск из app/:  uv run python scripts/import_gold.py [--check] [--manual PATH] [--out DIR]
"""
from __future__ import annotations

import argparse
import json
import re
import sys
import unicodedata
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT / "src") not in sys.path:
    sys.path.insert(0, str(ROOT / "src"))

import yaml

from dms_core.config import derived_dir, gold_xlsx, kit_dir
from dms_core.kit import nfc, scan

LABELS = ("allowed", "denied", "conditional", "insufficient")
CONFIDENCE = ("confirmed", "need_clarification", "contradictory", "insufficient")
SCENARIOS = ("sc01", "sc02_info", "sc02_clinic", "sc02_add_clinic")
DOC_TYPES = ("contract", "program", "clinics", "memo", "rules", "catalog", "amendment", "insurer_reply", "other")

# Заголовки листа «Кейсы» → ключи записи. Ищем по заголовку, а не по номеру колонки: порядок колонок может измениться.
COLUMNS = {
    "№": "id",
    "Сложность": "difficulty",
    "Группа компаний": "group",
    "Юрлицо / договор / программа": "entity",
    "СК": "insurer",
    "Сценарий": "scenario_raw",
    "Вывод (SC 01)": "label_raw",
    "Обращение (как от клиента)": "client_message",
    "Вопрос для ассистента": "question",
    "Эталонный ответ сопровождения": "reference_answer",
    "Документы и пункты": "gold_refs",
    "Что проверяет кейс": "checks",
    "Файлы в папке «Документы»": "files",
}

# Метки SC 01 по ТЗ п. 5.1 (план A5): уточнение в скобках → label_note. Порядок важен: «можно при условиях»
# проверяется раньше «можно».
_LABEL_WORDS = (("можно при условиях", "conditional"), ("недостаточно данных", "insufficient"),
                ("нельзя", "denied"), ("можно", "allowed"))
_CYR2LAT = str.maketrans({"К": "K"})


# ------------------------------------------------------------------ разбор xlsx

def parse_label(raw: str | None) -> tuple[str | None, str | None]:
    """«Нельзя (есть альтернатива)» → ("denied", "есть альтернатива"); «—» и пусто → (None, None)."""
    s = (raw or "").strip()
    if not s or s in {"—", "-", "–"}:
        return None, None
    m = re.match(r"^(.*?)\s*\((.+)\)\s*$", s)
    head, note = (m.group(1), m.group(2).strip()) if m else (s, None)
    head = head.strip().casefold()
    for words, label in _LABEL_WORDS:
        if head.startswith(words):
            return label, note
    raise ValueError(f"неизвестная метка SC 01: {s!r}")


def parse_date(text: str | None) -> str | None:
    """Первая дата «ДД.ММ.ГГГГ» в тексте обращения → ISO; нет даты → None."""
    m = re.search(r"\b(\d{2})\.(\d{2})\.(\d{4})\b", text or "")
    return date(int(m.group(3)), int(m.group(2)), int(m.group(1))).isoformat() if m else None


def read_cases(xlsx: Path) -> dict[str, dict]:
    import openpyxl

    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    rows = list(wb["Кейсы"].iter_rows(values_only=True))
    wb.close()
    head = [str(h).strip() if h is not None else "" for h in rows[0]]
    missing = [h for h in COLUMNS if h not in head]
    if missing:
        raise SystemExit(f"в листе «Кейсы» нет колонок: {missing}")
    idx = {COLUMNS[h]: head.index(h) for h in COLUMNS}
    out: dict[str, dict] = {}
    for r in rows[1:]:
        raw_id = r[idx["id"]]
        if not raw_id or not str(raw_id).strip().startswith(("К", "K")):
            continue
        rec = {k: (nfc(str(r[i])).strip() if r[i] is not None else None) for k, i in idx.items()}
        rec["id"] = rec["id"].translate(_CYR2LAT)
        out[rec["id"]] = rec
    return out


# ------------------------------------------------------------------ документы комплекта и проверка ссылок

_DASHES = dict.fromkeys(map(ord, "‐‑‒–—−"), "-")
_QUOTES = dict.fromkeys(map(ord, "«»“”„‟″"), '"')


def norm(s: str) -> str:
    """Сравнение цитат: NFC, ё→е, тире → «-», кавычки → «"», без меток сносок [^N], пробелы схлопнуты, casefold."""
    s = unicodedata.normalize("NFC", s).replace("ё", "е").replace("Ё", "Е")
    s = s.translate(_DASHES).translate(_QUOTES)
    s = re.sub(r"\[\^\d+\]", "", s)
    s = re.sub(r"\s+", " ", s.replace(" ", " "))
    return s.strip().casefold()


class Kit:
    """Доступ к блокам документов комплекта по (договор, тип документа, адрес блока). Документы читаются один раз."""

    def __init__(self, root: Path):
        self.root = root
        self.files: dict[tuple[str, str], Path] = {}
        for g in scan(root):
            for c in g.contracts:
                for f in c.files:
                    self.files.setdefault((c.id, f.doc_type), root / f.path)
        self._docx: dict[Path, object] = {}
        self._xlsx: dict[Path, tuple[str, list[list[str]]]] = {}

    def path(self, contract_id: str, doc_type: str) -> Path:
        p = self.files.get((contract_id, doc_type))
        if p is None:
            raise KeyError(f"{contract_id}.{doc_type}: документа нет в комплекте")
        return p

    def _read_docx(self, p: Path):
        if p not in self._docx:
            from dms_core.io.docx_reader import read_docx

            self._docx[p] = read_docx(p)
        return self._docx[p]

    def _read_xlsx(self, p: Path) -> tuple[str, list[list[str]]]:
        if p not in self._xlsx:
            import openpyxl

            wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
            ws = wb.worksheets[0]
            rows = [["" if v is None else str(v) for v in row] for row in ws.iter_rows(values_only=True)]
            self._xlsx[p] = (ws.title, rows)
            wb.close()
        return self._xlsx[p]

    def block(self, contract_id: str, doc_type: str, block: str) -> tuple[str, str | None, str]:
        """(текст блока, номер автонумерации или None, ссылка вида «c01.contract#p40» / «…#t6r4» / «…#Лист!569»)."""
        p = self.path(contract_id, doc_type)
        doc_id = f"{contract_id}.{doc_type}"
        if p.suffix.lower() == ".xlsx":
            sheet, rows = self._read_xlsx(p)
            n = int(block)
            if not 1 <= n <= len(rows):
                raise KeyError(f"{doc_id}: нет строки {n}")
            return " ".join(rows[n - 1]), None, f"{doc_id}#{sheet}!{n}"
        d = self._read_docx(p)
        m = re.fullmatch(r"p(\d+)", block)
        if m:
            para = d.paragraphs[int(m.group(1))]
            return " ".join([para.text, *para.footnotes]), para.number, f"{doc_id}#{block}"
        m = re.fullmatch(r"t(\d+)r(\d+)", block)
        if m:
            row = d.tables[int(m.group(1))].rows[int(m.group(2))]
            return " ".join(c.text for c in row), None, f"{doc_id}#{block}"
        raise KeyError(f"{doc_id}: непонятный адрес блока {block!r}")


def _num(s: str | None) -> str:
    return (s or "").strip().rstrip(".")


def verify_citation(kit: Kit, c: dict) -> tuple[str | None, list[str]]:
    """Проверка одной ссылки разметки по документу. Возвращает (адрес блока, список проблем без цитат)."""
    problems: list[str] = []
    try:
        text, number, ref = kit.block(c["contract_id"], c["doc_type"], str(c["block"]))
    except (KeyError, IndexError, ValueError) as e:
        return None, [f"блок не найден: {e}"]
    t = norm(text)
    q = norm(c["quote_hint"])
    pos = t.find(q)
    if pos < 0:
        problems.append("quote_hint не найден в блоке")
    for k, extra in enumerate(c.get("also") or []):
        if norm(extra) not in t:
            problems.append(f"also[{k}] не найден в блоке")
    clause = c.get("clause")
    if number and not re.search(r"[0-9IVXivx]", number):
        number = None          # маркер списка («-», «•», мягкий перенос) — не номер пункта
    if clause and number:
        # Автономер абзаца (docx_reader) обязан совпасть с clause: «2.6.1.» ↔ «2.6.1».
        if _num(number) != _num(clause):
            problems.append(f"номер абзаца {_num(number)!r} ≠ clause {clause!r}")
    elif clause and pos >= 0 and str(c["block"]).startswith("p") and not c.get("subitem"):
        # Номер набран текстом («2.3.\t…», «VII. …») — он в начале абзаца или (абзац-продолжение) у предыдущего
        # номерного: тогда проверяется только, что номер не противоречит началу текста.
        head = t[:12]
        if re.match(r"^[\divx]+(\.[\d]+)*\.?\s", head) and not head.startswith(norm(clause)):
            problems.append(f"номер в тексте не совпадает с clause {clause!r}")
    elif clause and pos >= 0 and c.get("path") and re.fullmatch(r"t\d+r\d+", str(c["block"])):
        # Программа Ренессанса: весь текст в одной ячейке, номер «1.4.» стоит прямо перед цитатой
        # (а не «1.3.» и не «2.1.4.»: слева от номера — начало, пробел или «/» разделителя строк).
        before = t[max(0, pos - 24):pos]
        if not re.search(rf"(^|[\s/]){re.escape(norm(clause))}\.\s*$", before):
            problems.append(f"перед цитатой нет номера {clause!r}")
    return ref, problems


# ------------------------------------------------------------------ сборка записей

_CIT_KEYS = ("contract_id", "doc_type", "clause", "clause_gold", "subitem", "path", "program", "quote_hint")


def build(cases: dict[str, dict], manual: dict, kit: Kit) -> tuple[list[dict], dict, list[str]]:
    records: list[dict] = []
    acceptance: dict[str, list[dict]] = {}
    errors: list[str] = []
    for cid in sorted(manual["cases"], key=lambda x: int(x[1:])):
        m = manual["cases"][cid]
        x = cases.get(cid)
        if x is None:
            errors.append(f"{cid}: нет в xlsx")
            continue
        label, note = parse_label(x["label_raw"])
        ctx = dict(m["context"])
        if ctx.get("date") is None:
            ctx["date"] = None
        req, acc = [], []
        for k, c in enumerate(m["required_citations"]):
            ref, problems = verify_citation(kit, c)
            for p in problems:
                errors.append(f"{cid} ссылка {k} ({c['contract_id']}.{c['doc_type']} {c['block']}): {p}")
            cit = {key: c.get(key) for key in _CIT_KEYS}
            is_table = ref is not None and not re.search(r"#p\d+$", ref)
            cit["table_ref"] = ref if is_table else None
            cit["clause_id"] = ref if (ref is not None and not is_table) else None
            if c.get("required", True):
                req.append(cit)
            acc.append({**cit, "also": list(c.get("also") or []), "doc_id": f"{c['contract_id']}.{c['doc_type']}",
                        "ref": ref,
                        "required": bool(c.get("required", True)), "verified": not problems})
        contracts = sorted({c["contract_id"] for c in req})
        rec = {
            "id": cid,
            "source": "organizers",
            "scenario": m["scenario"],
            "context": ctx,
            "question": x["question"],
            "client_message": x["client_message"],
            "difficulty": x["difficulty"],
            "insurer": x["insurer"],
            "contracts": contracts,
            "expected": {
                "label": label,
                "label_note": note,
                "label_raw": x["label_raw"],
                "confidence": m["confidence"],
                "key_facts": list(m["key_facts"]),
                "required_citations": req,
                "must_not": list(m.get("must_not") or []),
                "questions": list(m.get("questions") or []),
                "automation": m.get("automation"),
            },
            "reference_answer": x["reference_answer"],
            "gold_refs": x["gold_refs"],
            "checks": x["checks"],
            "trap": m.get("trap"),
            "notes": m.get("notes"),
        }
        if m.get("per_contract"):
            rec["expected"]["per_contract"] = m["per_contract"]
        # Дата из обращения должна совпасть с датой контекста, если она там есть (К10: 19.10.2026).
        d = parse_date(x["client_message"])
        if d and ctx.get("date") and d != ctx["date"]:
            errors.append(f"{cid}: дата обращения {d} ≠ context.date {ctx['date']}")
        records.append(rec)
        acceptance[cid] = acc
    missing = sorted(set(cases) - set(manual["cases"]))
    if missing:
        errors.append(f"нет ручной разметки для: {missing}")
    acc_doc = {"version": 1, "source": "scripts/import_gold.py", "note":
               "Пункты/строки, которые инструменты обязаны найти (required=false — полезные, но не обязательные).",
               "cases": acceptance}
    return records, acc_doc, errors


def summary(records: list[dict]) -> str:
    lines = ["id   сценарий   договоры              метка (уточнение)                         факты  ссылки"]
    for r in records:
        e = r["expected"]
        lab = (e["label"] or "—") + (f" ({e['label_note']})" if e["label_note"] else "")
        lines.append(f"{r['id']:<4} {r['scenario']:<10} {','.join(r['contracts']):<21} {lab:<41} "
                     f"{len(e['key_facts']):>5}  {len(e['required_citations']):>6}")
    return "\n".join(lines)


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--xlsx", type=Path, default=None, help="файл кейсов (по умолчанию config.gold_xlsx())")
    ap.add_argument("--manual", type=Path, default=None, help="ручная разметка (data/derived/gold/k10_k19_manual.yaml)")
    ap.add_argument("--out", type=Path, default=None, help="папка результата (data/derived/gold/)")
    ap.add_argument("--check", action="store_true", help="только проверить ссылки, файлы не писать")
    a = ap.parse_args(argv)
    out = a.out or derived_dir() / "gold"
    xlsx = a.xlsx or gold_xlsx()
    manual_path = a.manual or out / "k10_k19_manual.yaml"
    for p, what in ((xlsx, "кейсы организаторов"), (manual_path, "ручная разметка"), (kit_dir(), "комплект")):
        if not p.exists():
            print(f"нет файла ({what}): {p}", file=sys.stderr)
            return 2
    manual = yaml.safe_load(manual_path.read_text(encoding="utf-8"))
    records, acc, errors = build(read_cases(xlsx), manual, Kit(kit_dir()))
    print(summary(records))
    n_acc = sum(len(v) for v in acc["cases"].values())
    print(f"\nкейсов: {len(records)}; ссылок для приёмки: {n_acc}; ошибок проверки: {len(errors)}")
    for e in errors:
        print("  ОШИБКА", e)
    if errors:
        return 1
    if a.check:
        return 0
    out.mkdir(parents=True, exist_ok=True)
    with (out / "k10_k19.jsonl").open("w", encoding="utf-8") as fh:
        for r in records:
            fh.write(json.dumps(r, ensure_ascii=False) + "\n")
    (out / "acceptance_k10_k19.json").write_text(json.dumps(acc, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"записано: {out / 'k10_k19.jsonl'}, {out / 'acceptance_k10_k19.json'}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
